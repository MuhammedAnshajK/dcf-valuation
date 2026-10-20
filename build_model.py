"""Builds nike_valuation.xlsx: a DCF (perpetuity growth and exit multiple) and
trading-comparables valuation of NIKE, Inc. using the base case of the Nike
five-year forecast (three-statement-model project) and market data at
19 Oct 2026. Run: python build_model.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

wb = Workbook(); F = "Arial"
def fnt(color="000000", bold=False, italic=False, size=10): return Font(name=F, color=color, bold=bold, italic=italic, size=size)
BLUE, BLACK, GREEN, GREY = "0000FF", "000000", "008000", "595959"
HFILL = PatternFill("solid", fgColor="1F4E78"); YEL = PatternFill("solid", fgColor="FFFF00")
LIGHT = PatternFill("solid", fgColor="DDEBF7")
TOP = Border(top=Side(style="thin")); DBL = Border(top=Side(style="thin"), bottom=Side(style="double"))
NUM = '#,##0;(#,##0);"-"'; PCT = '0.0%;(0.0%);"-"'; MULT = '0.0"x"'; PS = '$#,##0.00;($#,##0.00);"-"'; DEC = '0.00'

def cell(ws, ref, v, fmt=None, color=None, bold=False, fill=None, border=None, italic=False):
    c = ws[ref]; c.value = v
    if color is None:
        color = (GREEN if "!" in v else BLACK) if isinstance(v, str) and v.startswith("=") else (BLUE if isinstance(v, (int, float)) else BLACK)
    c.font = fnt(color, bold, italic)
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if border: c.border = border
    return c

def title(ws, t, sub="NIKE, Inc. | $ millions unless stated | market data at 19 Oct 2026"):
    ws["B1"] = t; ws["B1"].font = fnt(bold=True, size=14)
    ws["B2"] = sub; ws["B2"].font = fnt(GREY, italic=True)
    ws.column_dimensions["A"].width = 2; ws.column_dimensions["B"].width = 40
    ws.sheet_view.showGridLines = False

def header(ws, r, labels, start_col=3):
    ws[f"B{r}"].fill = HFILL
    for i, lab in enumerate(labels):
        col = chr(64 + start_col + i); c = ws[f"{col}{r}"]; c.value = lab
        c.font = fnt("FFFFFF", True); c.fill = HFILL; c.alignment = Alignment(horizontal="right")
        ws.column_dimensions[col].width = 12

YRS = ["FY2026A", "FY2027E", "FY2028E", "FY2029E", "FY2030E", "FY2031E"]; YC = list("CDEFGH"); PC = YC[1:]
def prv(c): return chr(ord(c) - 1)

# ---------------- Cover ----------------
cv = wb.active; cv.title = "Cover"; cv.sheet_view.showGridLines = False
cv.column_dimensions["A"].width = 2; cv.column_dimensions["B"].width = 22; cv.column_dimensions["C"].width = 95
cv["B2"] = "NIKE, Inc.: DCF & Comparable Companies Valuation"; cv["B2"].font = fnt(bold=True, size=18)
cv["B3"] = "Valuation date 19 Oct 2026 | share price $37.21 | $ millions unless stated"; cv["B3"].font = fnt(GREY, italic=True)
info = [
 ("Purpose", "Values Nike three ways (DCF with perpetuity growth, DCF with an exit multiple, and trading comparables) and compares them with the share price on one football field chart."),
 ("Data", "Cash flows from the base case of the Nike five-year forecast (github.com/MuhammedAnshajK/three-statement-model), whose FY2027 matches Nike's guidance of 1 Oct 2026. Share price, beta and peer multiples from stockanalysis.com at 19 Oct 2026; 10-year Treasury yield from YCharts (16 Oct 2026). Net debt from the 31 May 2026 balance sheet (10-K)."),
 ("How to use", "Change the blue or yellow cells on the WACC, DCF and Comps sheets. Everything else is a formula."),
 ("Sheets", "Summary > Forecast > WACC > DCF > Sensitivity > Comps"),
 ("Colour key", "Blue = hardcoded input | Black = formula | Green = link to another sheet | Yellow fill = key assumption"),
 ("Note", "A learning project, not investment advice."),
 ("Author", "Muhammed Anshaj K"),
]
for i, (k, v) in enumerate(info):
    cv[f"B{5+i}"] = k; cv[f"B{5+i}"].font = fnt(bold=True); cv[f"B{5+i}"].alignment = Alignment(vertical="top")
    cv[f"C{5+i}"] = v; cv[f"C{5+i}"].font = fnt(); cv[f"C{5+i}"].alignment = Alignment(wrap_text=True, vertical="top")

# ---------------- Forecast (from three-statement model) ----------------
fc = wb.create_sheet("Forecast"); title(fc, "Forecast inputs"); header(fc, 4, YRS)
data = {
 5: ("Revenue", [46398, 42441.5, 43237.1, 45140.2, 47196.8, 49130.0]),
 6: ("EBITDA", [4544, 2673.9, 3631.9, 4604.3, 5474.8, 6239.5]),
 7: ("Depreciation & amortisation", [747, 679.1, 691.8, 722.2, 755.1, 786.1]),
 9: ("Capital expenditure", [684, 679.1, 691.8, 722.2, 755.1, 786.1]),
 10: ("Operating working capital (AR + inventory + prepaid - AP - accrued)", [5884, 4851.7, 4541.0, 4324.3, 4419.1, 4569.8]),
 11: ("Diluted EPS ($)", [2.10, 1.09, 1.63, 2.19, 2.73, 3.23]),
}
for r, (lab, vals) in data.items():
    cell(fc, f"B{r}", lab)
    for c, v in zip(YC, vals):
        if v is not None: cell(fc, f"{c}{r}", v, PS if r == 11 else NUM)
cell(fc, "B8", "EBIT", bold=True)
for c in YC: cell(fc, f"{c}8", f"={c}6-{c}7", NUM, bold=True, border=TOP)
fc.column_dimensions["B"].width = 58
cell(fc, "B13", "Balance sheet at 31 May 2026 and market data", bold=True)
cell(fc, "B14", "Cash and short-term investments"); cell(fc, "C14", 9027, NUM)
cell(fc, "B15", "Debt (excluding leases)"); cell(fc, "C15", 7942, NUM)
cell(fc, "B16", "Shares outstanding (m)"); cell(fc, "C16", 1485, NUM, fill=YEL)
cell(fc, "B17", "Share price at 19 Oct 2026 ($)"); cell(fc, "C17", 37.21, PS, fill=YEL)
cell(fc, "B18", "Market capitalisation"); cell(fc, "C18", "=C16*C17", NUM)
cell(fc, "B19", "52-week low / high ($)"); cell(fc, "C19", 31.97, PS); cell(fc, "D19", 70.44, PS)
cell(fc, "B21", "Revenue, EBITDA, D&A, capex, working capital and EPS: base case of the Nike forecast model (rounded to 0.1). FY2026A EPS is reported diluted EPS.", italic=True, color=GREY)
cell(fc, "B22", "Shares: market cap $55.27bn / $37.21 (stockanalysis.com). Debt and cash: Nike 10-K, 31 May 2026.", italic=True, color=GREY)

# ---------------- WACC ----------------
w = wb.create_sheet("WACC"); title(w, "Discount rate (WACC)"); w.column_dimensions["C"].width = 12; w.column_dimensions["D"].width = 60
rows = [
 (5, "Risk-free rate", 0.0509, PCT, "US 10-year Treasury yield, 16 Oct 2026 (YCharts)"),
 (6, "Equity risk premium", 0.05, PCT, "Assumption, within the common 4% to 6% range"),
 (7, "Beta (levered)", 1.10, DEC, "Nike 5-year beta (stockanalysis.com)"),
 (8, "Cost of equity (CAPM)", "=C5+C7*C6", PCT, "Risk-free + beta x equity risk premium"),
 (10, "Pre-tax cost of debt", 0.059, PCT, "Risk-free rate plus about 0.8% for an A-rated issuer"),
 (11, "Tax rate", 0.19, PCT, "Matches the forecast tax rate"),
 (12, "After-tax cost of debt", "=C10*(1-C11)", PCT, ""),
 (14, "Debt / (debt + equity)", "=Forecast!C15/(Forecast!C15+Forecast!C18)", PCT, "Market-value weights: debt $7.9bn, market cap $55.3bn"),
 (15, "Equity / (debt + equity)", "=1-C14", PCT, ""),
 (17, "WACC", "=C15*C8+C14*C12", PCT, "Weighted average cost of capital"),
]
for r, lab, v, fmt, note in rows:
    cell(w, f"B{r}", lab, bold=(r in (8, 17)))
    cell(w, f"C{r}", v, fmt, bold=(r in (8, 17)), fill=(YEL if r in (5, 6, 7) else (LIGHT if r == 17 else None)),
         border=(DBL if r == 17 else None))
    cell(w, f"D{r}", note, italic=True, color=GREY)

# ---------------- DCF ----------------
d = wb.create_sheet("DCF"); title(d, "Discounted cash flow"); header(d, 4, YRS)
cell(d, "B5", "EBIT")
for c in YC: cell(d, f"{c}5", f"=Forecast!{c}8", NUM)
cell(d, "B6", "Less: tax on EBIT")
for c in PC: cell(d, f"{c}6", f"=-{c}5*WACC!$C$11", NUM)
cell(d, "B7", "NOPAT", bold=True)
for c in PC: cell(d, f"{c}7", f"={c}5+{c}6", NUM, bold=True, border=TOP)
cell(d, "B8", "Add: depreciation & amortisation")
for c in PC: cell(d, f"{c}8", f"=Forecast!{c}7", NUM)
cell(d, "B9", "Less: capital expenditure")
for c in PC: cell(d, f"{c}9", f"=-Forecast!{c}9", NUM)
cell(d, "B10", "Less: increase in working capital")
for c in PC: cell(d, f"{c}10", f"=-(Forecast!{c}10-Forecast!{prv(c)}10)", NUM)
cell(d, "B11", "Unlevered free cash flow", bold=True)
for c in PC: cell(d, f"{c}11", f"=SUM({c}7:{c}10)", NUM, bold=True, border=TOP)
cell(d, "B13", "Discount period (years)")
for i, c in enumerate(PC): cell(d, f"{c}13", round(0.61 + i, 2), '0.00')
cell(d, "B14", "Discount factor")
for c in PC: cell(d, f"{c}14", f"=1/(1+$C$18)^{c}13", '0.000')
cell(d, "B15", "Present value of free cash flow")
for c in PC: cell(d, f"{c}15", f"={c}11*{c}14", NUM)

cell(d, "B17", "Assumptions", bold=True)
cell(d, "B18", "WACC"); cell(d, "C18", "=WACC!C17", PCT)
cell(d, "B19", "Terminal growth rate"); cell(d, "C19", 0.03, PCT, fill=YEL)
cell(d, "B20", "Exit EV / EBITDA multiple"); cell(d, "C20", 10.0, MULT, fill=YEL)

# Perpetuity growth method
cell(d, "B22", "Valuation", bold=True)
cell(d, "C22", "Perpetuity", bold=True); cell(d, "E22", "Exit multiple", bold=True)
for ref in ("C22", "E22"): d[ref].alignment = Alignment(horizontal="right")
rowsP = [
 (23, "Sum of PV of free cash flows", "=SUM(D15:H15)", "=SUM(D15:H15)"),
 (24, "Terminal value (FY2031)", "=H11*(1+C19)/(C18-C19)", "=Forecast!H6*C20"),
 (25, "PV of terminal value", "=C24*H14", "=E24*H14"),
 (26, "Enterprise value", "=C23+C25", "=E23+E25"),
 (27, "Less: debt", "=-Forecast!C15", "=-Forecast!C15"),
 (28, "Add: cash", "=Forecast!C14", "=Forecast!C14"),
 (29, "Equity value", "=SUM(C26:C28)", "=SUM(E26:E28)"),
 (30, "Value per share ($)", "=C29/Forecast!C16", "=E29/Forecast!C16"),
 (31, "Upside / (downside) to share price", "=C30/Forecast!C17-1", "=E30/Forecast!C17-1"),
 (32, "Terminal value as % of EV", "=C25/C26", "=E25/E26"),
 (33, "Implied exit multiple / implied growth", "=C24/Forecast!H6", "=(E24*C18-H11)/(E24+H11)"),
]
for r, lab, p, e in rowsP:
    bold = r in (26, 29, 30)
    fmt = PS if r == 30 else (PCT if r in (31, 32) else NUM)
    cell(d, f"B{r}", lab, bold=bold)
    cell(d, f"C{r}", p, MULT if r == 33 else fmt, bold=bold, border=(TOP if r in (26, 29) else None))
    cell(d, f"E{r}", e, PCT if r == 33 else fmt, bold=bold, border=(TOP if r in (26, 29) else None))
for r in (30,): d[f"C{r}"].fill = LIGHT; d[f"E{r}"].fill = LIGHT
for i, t in enumerate(["Cash flows are discounted to 19 Oct 2026 at each 31 May year end (0.61 years to the FY2027 year end).",
                       "FY2027 is counted in full, a small simplification. Share-based pay is treated as a real cost (not added back).",
                       "Terminal value uses FY2031 free cash flow grown one year."]):
    cell(d, f"B{35 + i}", t, italic=True, color=GREY)

# ---------------- Sensitivity ----------------
s = wb.create_sheet("Sensitivity"); title(s, "Sensitivity: value per share ($)", "Perpetuity growth method | WACC down the side, terminal growth across the top | share price $37.21")
s.column_dimensions["B"].width = 18
waccs = [-0.01, -0.005, 0, 0.005, 0.01]; gs = [-0.01, -0.005, 0, 0.005, 0.01]
cell(s, "B4", "WACC \\ growth", bold=True)
for j, dg in enumerate(gs):
    col = chr(ord("C") + j); s.column_dimensions[col].width = 12
    cell(s, f"{col}4", f"=DCF!$C$19+({dg})", PCT, bold=True, fill=LIGHT)
for i, dw in enumerate(waccs):
    r = 5 + i
    cell(s, f"B{r}", f"=DCF!$C$18+({dw})", PCT, bold=True, fill=LIGHT)
    for j in range(5):
        col = chr(ord("C") + j)
        f = (f"=(SUMPRODUCT(DCF!$D$11:$H$11/(1+$B{r})^DCF!$D$13:$H$13)"
             f"+DCF!$H$11*(1+{col}$4)/($B{r}-{col}$4)/(1+$B{r})^DCF!$H$13"
             f"-Forecast!$C$15+Forecast!$C$14)/Forecast!$C$16")
        cell(s, f"{col}{r}", f, PS)
s["E7"].fill = YEL
cell(s, "B11", "Centre cell (yellow) equals the base case on the DCF sheet.", italic=True, color=GREY)
cell(s, "B12", "Check: centre less DCF value per share"); cell(s, "E12", "=ROUND(E7-DCF!C30,6)", PS)

# ---------------- Comps ----------------
cp = wb.create_sheet("Comps"); title(cp, "Trading comparables", "Athletic footwear and apparel peers | stockanalysis.com, 19 Oct 2026 (Puma 7 Oct) | n.m. = not meaningful")
cp.column_dimensions["B"].width = 30
header(cp, 4, ["EV / Revenue", "EV / EBITDA", "Forward P / E"])
NM = "n.m."
peers = [("adidas (Germany)", 1.24, 9.77, 15.21), ("Puma (Germany)", 0.79, NM, NM),
         ("Lululemon (US)", 1.04, 4.86, 11.73), ("Deckers Outdoor (US)", 1.95, 8.13, 11.43),
         ("On Holding (Switzerland)", 2.78, 18.25, 19.02), ("Under Armour (US)", 0.61, 17.44, NM)]
for i, (n, a_, b_, c_) in enumerate(peers):
    r = 5 + i; cell(cp, f"B{r}", n)
    for col, v in zip("CDE", (a_, b_, c_)):
        c = cell(cp, f"{col}{r}", v, MULT)
        if v == NM: c.alignment = Alignment(horizontal="right")
stats = [(12, "25th percentile", "QUARTILE({c}5:{c}10,1)"), (13, "Median", "MEDIAN({c}5:{c}10)"), (14, "75th percentile", "QUARTILE({c}5:{c}10,3)")]
for r, lab, f in stats:
    cell(cp, f"B{r}", lab, bold=True)
    for col in "CDE": cell(cp, f"{col}{r}", "=" + f.format(c=col), MULT, bold=True, border=(TOP if r == 12 else None))
cell(cp, "B15", "Nike today (for reference)", italic=True, color=GREY)
for col, v in zip("CDE", (1.26, 11.86, 28.41)): cell(cp, f"{col}15", v, MULT, italic=True)
cell(cp, "B16", "Nike metric", bold=True)
cell(cp, "C16", 45890, NUM); cell(cp, "D16", 4890, NUM); cell(cp, "E16", 1.25, PS, fill=YEL)
cell(cp, "B17", "Revenue and EBITDA: last twelve months to Aug 2026 (stockanalysis.com). EPS: midpoint of FY2027 guidance ($1.15 to $1.35, adjusted).", italic=True, color=GREY)
cell(cp, "B18", "Implied value per share ($)", bold=True)
for r, lab in ((19, "At 25th percentile"), (20, "At median"), (21, "At 75th percentile")):
    cell(cp, f"B{r}", lab); sr = r - 7
    for col in "CD":  # EV multiples: add net cash
        cell(cp, f"{col}{r}", f"=({col}{sr}*{col}$16-Forecast!$C$15+Forecast!$C$14)/Forecast!$C$16", PS)
    cell(cp, f"E{r}", f"=E{sr}*E$16", PS)
cell(cp, "B23", "Puma's EBITDA is negative and its earnings and Under Armour's are loss-making or near zero, so those multiples are excluded (n.m.).", italic=True, color=GREY)

# ---------------- Summary / football field ----------------
sm = wb.create_sheet("Summary"); title(sm, "Valuation summary", "Value per share ($) | the bars in the chart run from low to high")
sm.column_dimensions["B"].width = 34
header(sm, 4, ["Low", "High", "Range", "Base"])
summ = [
 (5, "DCF: perpetuity growth (sensitivity)", "=MIN(Sensitivity!C5:G9)", "=MAX(Sensitivity!C5:G9)", "=DCF!C30"),
 (6, "DCF: exit multiple (9.0x to 11.0x)", "=(DCF!C23+(DCF!C20-1)*Forecast!H6*DCF!H14-Forecast!C15+Forecast!C14)/Forecast!C16",
     "=(DCF!C23+(DCF!C20+1)*Forecast!H6*DCF!H14-Forecast!C15+Forecast!C14)/Forecast!C16", "=DCF!E30"),
 (7, "Comps: EV / EBITDA (25th to 75th pct)", "=Comps!D19", "=Comps!D21", "=Comps!D20"),
 (8, "Comps: EV / Revenue (25th to 75th pct)", "=Comps!C19", "=Comps!C21", "=Comps!C20"),
 (9, "Comps: forward P / E (25th to 75th pct)", "=Comps!E19", "=Comps!E21", "=Comps!E20"),
 (10, "52-week trading range", "=Forecast!C19", "=Forecast!D19", "=Forecast!C17"),
]
for r, lab, lo, hi, base in summ:
    cell(sm, f"B{r}", lab); cell(sm, f"C{r}", lo, PS); cell(sm, f"D{r}", hi, PS)
    cell(sm, f"E{r}", f"=D{r}-C{r}", PS); cell(sm, f"F{r}", base, PS)
cell(sm, "B12", "Current share price (19 Oct 2026)", bold=True); cell(sm, "F12", "=Forecast!C17", PS, bold=True)
cell(sm, "B13", "DCF midpoint (perpetuity and exit multiple)", bold=True)
cell(sm, "F13", "=AVERAGE(F5:F6)", PS, bold=True, fill=LIGHT)
cell(sm, "B14", "Upside / (downside) to share price", bold=True); cell(sm, "F14", "=F13/F12-1", PCT, bold=True)
cell(sm, "B16", "Base column: DCF base cases, comps at the median, and the current price for the trading range.", italic=True, color=GREY)
cell(sm, "B17", "Forward P/E gives a low value because FY2027 earnings are temporarily depressed; the DCF captures the recovery.", italic=True, color=GREY)

ch = BarChart(); ch.type = "bar"; ch.grouping = "stacked"; ch.overlap = 100
ch.title = "Football field: value per share ($)"; ch.y_axis.title = "$ per share"
ch.add_data(Reference(sm, min_col=3, min_row=4, max_row=10), titles_from_data=True)
ch.add_data(Reference(sm, min_col=5, min_row=4, max_row=10), titles_from_data=True)
ch.set_categories(Reference(sm, min_col=2, min_row=5, max_row=10))
ch.series[0].graphicalProperties.noFill = True; ch.series[0].graphicalProperties.line.noFill = True
ch.series[1].graphicalProperties.solidFill = "1F4E78"
ch.legend = None; ch.x_axis.scaling.orientation = "maxMin"; ch.y_axis.majorGridlines = None
ch.x_axis.delete = False; ch.y_axis.delete = False; ch.y_axis.number_format = '$#,##0'
ch.width = 22; ch.height = 9
sm.add_chart(ch, "B19")

wb.move_sheet("Summary", offset=-(len(wb.sheetnames) - 2))
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
wb.save("nike_valuation.xlsx")
print("Saved nike_valuation.xlsx. Open it in Excel (or run a recalculation) to see values.")
